import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import json

class ExcelExporter:
    def __init__(self):
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = "Resume Data"
        
        # Define styles
        self.orange_header_font = Font(bold=True, color="FFFFFF", size=14)
        self.orange_header_fill = PatternFill(start_color="FF6600", end_color="FF6600", fill_type="solid")
        self.column_header_font = Font(bold=True, color="FFFFFF")  # Make header font white
        self.column_header_fill = PatternFill(start_color="FF6600", end_color="FF6600", fill_type="solid")  # Orange header
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        self.center_alignment = Alignment(horizontal='center', vertical='center')
        self.wrap_alignment = Alignment(wrap_text=True, vertical='top')
    
    def create_resume_sheet(self, parsed_data, filename="Resume"):
        """Create an Excel sheet with parsed resume data"""
        # Clear the sheet
        self.ws.delete_rows(1, self.ws.max_row)
        
        # Define column headers based on JSON structure
        headers = [
            "Name", "Email", "Phone", "GitHub", "LinkedIn", "School", "College", 
            "10th Marks (%)", "12th Marks (%)", "CGPA", "Internships", 
            "Internship Projects", "Personal Projects", "Tech Stack", "Achievements", "Languages"
        ]
        
        # No merged title/header row
        # Add column headers in row 1 (was row 2)
        for col, header in enumerate(headers, 1):
            cell = self.ws.cell(row=1, column=col, value=header)
            cell.font = self.column_header_font
            cell.fill = self.column_header_fill
            cell.alignment = self.center_alignment
            cell.border = self.border
        
        # Add data in row 2 (was row 3)
        data_values = [
            parsed_data.get('name', 'Not Specified'),
            parsed_data.get('email', 'Not Specified'),
            parsed_data.get('phone', 'Not Specified'),
            parsed_data.get('github', 'Not Specified'),
            parsed_data.get('linkedin', 'Not Specified'),
            parsed_data.get('school', 'Not Specified'),
            parsed_data.get('college', 'Not Specified'),
            parsed_data.get('tenth_marks', 'Not Specified'),
            parsed_data.get('twelfth_marks', 'Not Specified'),
            parsed_data.get('cgpa', 'Not Specified'),
            self._format_list(parsed_data.get('internships', [])),
            self._format_list(parsed_data.get('internship_projects', [])),
            self._format_list(parsed_data.get('personal_projects', [])),
            self._format_list(parsed_data.get('tech_stack', [])),
            self._format_list(parsed_data.get('achievements', [])),
            self._format_list(parsed_data.get('languages', []))
        ]
        
        for col, value in enumerate(data_values, 1):
            cell = self.ws.cell(row=2, column=col, value=str(value) if value else 'Not Specified')
            cell.border = self.border
            cell.alignment = self.wrap_alignment
        
        self.ws.row_dimensions[2].height = 60
        self._adjust_column_widths()
        return self.wb
    
    def create_multiple_resumes_sheet(self, multiple_resumes_data):
        """Create an Excel sheet with multiple resume data"""
        # Clear the sheet
        self.ws.delete_rows(1, self.ws.max_row)
        
        # Define column headers
        headers = [
            "S.No", "File Name", "Name", "Email", "Phone", "GitHub", "LinkedIn", 
            "School", "College", "10th Marks (%)", "12th Marks (%)", "CGPA", 
            "Internships", "Internship Projects", "Personal Projects", 
            "Tech Stack", "Achievements", "Languages"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = self.ws.cell(row=1, column=col, value=header)
            cell.font = self.column_header_font
            cell.fill = self.column_header_fill
            cell.alignment = self.center_alignment
            cell.border = self.border
        
        # Add data rows starting from row 2 (was row 3)
        for row_idx, resume_data in enumerate(multiple_resumes_data, 2):
            data_values = [
                row_idx - 1,  # S.No
                resume_data.get('filename', 'Unknown'),
                resume_data.get('name', 'Not Specified'),
                resume_data.get('email', 'Not Specified'),
                resume_data.get('phone', 'Not Specified'),
                resume_data.get('github', 'Not Specified'),
                resume_data.get('linkedin', 'Not Specified'),
                resume_data.get('school', 'Not Specified'),
                resume_data.get('college', 'Not Specified'),
                resume_data.get('tenth_marks', 'Not Specified'),
                resume_data.get('twelfth_marks', 'Not Specified'),
                resume_data.get('cgpa', 'Not Specified'),
                self._format_list(resume_data.get('internships', [])),
                self._format_list(resume_data.get('internship_projects', [])),
                self._format_list(resume_data.get('personal_projects', [])),
                self._format_list(resume_data.get('tech_stack', [])),
                self._format_list(resume_data.get('achievements', [])),
                self._format_list(resume_data.get('languages', []))
            ]
            
            for col, value in enumerate(data_values, 1):
                cell = self.ws.cell(row=row_idx, column=col, value=str(value) if value else 'Not Specified')
                cell.border = self.border
                cell.alignment = self.wrap_alignment
                
                # Highlight error rows
                if resume_data.get('error'):
                    cell.fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        
        # Set row heights for better readability
        for row in range(2, len(multiple_resumes_data) + 2):
            self.ws.row_dimensions[row].height = 30
        
        self._adjust_column_widths()
        return self.wb
    
    def _format_list(self, items):
        """Format a list of items for display in Excel"""
        if not items:
            return "Not Specified"
        if isinstance(items, list):
            return ", ".join([str(item) for item in items])
        return str(items)
    
    def _adjust_column_widths(self):
        """Auto-adjust column widths based on content"""
        for column in self.ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            # Set minimum and maximum widths
            adjusted_width = min(max(max_length + 2, 12), 30)
            self.ws.column_dimensions[column_letter].width = adjusted_width
    
    def save_to_bytes(self):
        """Save the workbook to bytes for download"""
        buffer = BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

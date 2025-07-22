from pdfminer.high_level import extract_text
from pdfminer.layout import LAParams
import docx
from io import BytesIO
from .extractors.name_extractor import NameExtractor
from .extractors.details_extractor import DetailsExtractor
from .extractors.education_extractor import EducationExtractor
from .extractors.internship_extractor import InternshipExtractor
from .extractors.project_details import ProjectDetailsExtractor
from .extractors.achievements_extractor import AchievementsExtractor
from .extractors.linkedin_extractor import LinkedInExtractor
from .extractors.language_extractor import LanguageExtractor
from .utils.ocr_extractor import OCRExtractor


class ResumeParser:
    def __init__(self):
        self.name_extractor = NameExtractor()
        self.details_extractor = DetailsExtractor()
        self.education_extractor = EducationExtractor()
        self.internship_extractor = InternshipExtractor()
        self.project_extractor = ProjectDetailsExtractor()
        self.achievements_extractor = AchievementsExtractor()
        self.linkedin_extractor = LinkedInExtractor()
        self.language_extractor = LanguageExtractor()
        self.ocr_extractor = OCRExtractor()
    
    def extract_text_from_pdf(self, file):
        """Extract text from PDF file with OCR fallback"""
        try:
            file_bytes = file.read()
            file_size_mb = len(file_bytes) / (1024 * 1024)
            
            laparams = LAParams(
                boxes_flow=0.5,
                word_margin=0.1,
                char_margin=2.0,
                line_margin=0.5
            )
            
            text = extract_text(BytesIO(file_bytes), laparams=laparams)
            
            print(text)
            

            if self.ocr_extractor.is_image_heavy_pdf(text, file_size_mb):
                print("PDF appears to be image-heavy, applying OCR...")
                ocr_text = self.ocr_extractor.extract_text_from_pdf_images(file_bytes)
                
                if ocr_text.strip():
                    combined_text = text + "\n\n--- OCR EXTRACTED TEXT ---\n" + ocr_text
                    print(f"Combined text length: {len(combined_text)} characters")
                    return combined_text
            
            return text
            
        except Exception as e:
            try:
                print(f"PDFMiner extraction failed: {str(e)}, trying OCR...")
                file_bytes = file.read()
                ocr_text = self.ocr_extractor.extract_text_from_pdf_images(file_bytes)
                return ocr_text
            except Exception as ocr_error:
                raise Exception(f"Both PDFMiner and OCR extraction failed: {str(e)}, OCR: {str(ocr_error)}")
    
    def extract_text_from_docx(self, file):
        """Extract text from DOCX file using python-docx2txt for better extraction"""
        try:
            import docx2txt
          
            file.seek(0)
            file_bytes = file.read()

            if not file_bytes and hasattr(file, 'file'):
                file.file.seek(0)
                file_bytes = file.file.read()
            import tempfile
            with tempfile.NamedTemporaryFile(delete=True, suffix=".docx") as tmp:
                tmp.write(file_bytes)
                tmp.flush()
                text = docx2txt.process(tmp.name)
            if not text.strip():
                from docx import Document
                doc = Document(BytesIO(file_bytes))
                for paragraph in doc.paragraphs:
                    text += paragraph.text + "\n"
                if not text.strip():
                    for table in doc.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                cell_text = cell.text.strip()
                                if cell_text:
                                    text += cell_text + "\n"
            return text
        except Exception as e:
            raise Exception(f"Error reading DOCX: {str(e)}")
    
    def parse_resume(self, file):
        """Main method to parse resume and extract all information"""
        try:
            if file.name.lower().endswith('.pdf'):
                text = self.extract_text_from_pdf(file)
            elif file.name.lower().endswith(('.docx', '.doc')):
                text = self.extract_text_from_docx(file)
            elif file.name.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp', '.tiff')):
             
                file_bytes = file.read()
                text = self.ocr_extractor.extract_text_from_image(file_bytes)
            else:
                raise Exception("Unsupported file format")
            
            if not text.strip():
                raise Exception("No text could be extracted from the file")
            
            print(f"Final extracted text length: {len(text)} characters")
            
            try:
                with open('/Users/happyfox/Desktop/second/extracted_resumes.txt', 'a', encoding='utf-8') as f:
                    f.write(f"\n\n--- Resume: {file.name} ---\n")
                    f.write(text)
                    f.write("\n\n")
            except Exception as file_write_error:
                print(f"Warning: Could not write extracted text to file: {file_write_error}")
            
            parsed_data = {
                'name': self.name_extractor.extract_name(text, filename=file.name),
                'email': self.details_extractor.extract_email(text),
                'phone': self.details_extractor.extract_phone(text),
                'github': self.details_extractor.extract_github(text),
                'linkedin': self.linkedin_extractor.extract_linkedin_url(text),
                'school': self.education_extractor.extract_school(text),
                'college': self.education_extractor.extract_college(text),
                'tenth_marks': self.education_extractor.extract_tenth_marks(text),
                'twelfth_marks': self.education_extractor.extract_twelfth_marks(text),
                'cgpa': self.education_extractor.extract_cgpa(text),
                'internships': self.internship_extractor.extract_internships(text),
                'internship_projects': self.internship_extractor.extract_internship_projects(text),
                'personal_projects': self.project_extractor.extract_personal_projects(text),
                'achievements': self.achievements_extractor.extract_achievements(text),
                'tech_stack': self._extract_tech_stack(text),
                'languages': self.language_extractor.extract_languages(text)
            }
            
            return parsed_data
            
        except Exception as e:
            raise Exception(f"Error parsing resume: {str(e)}")
    
    def _extract_tech_stack(self, text):
        """Extract technical skills/tech stack from resume text"""
      
        technologies = [
            'Python', 'Java', 'JavaScript', 'C++', 'C#', 'C', 'PHP', 'Ruby', 'Go', 'Rust',
            'React', 'Angular', 'Vue', 'Node.js', 'Express', 'Django', 'Flask', 'Spring',
            'HTML', 'CSS', 'Bootstrap', 'Tailwind', 'SASS', 'SCSS',
            'MySQL', 'PostgreSQL', 'MongoDB', 'Redis', 'SQLite', 'Oracle',
            'AWS', 'Azure', 'GCP', 'Docker', 'Kubernetes', 'Git', 'Jenkins',
            'TensorFlow', 'PyTorch', 'Scikit-learn', 'Pandas', 'NumPy',
            'Flutter', 'React Native', 'Android', 'iOS', 'Swift', 'Kotlin'
        ]
        
        found_tech = []
        text_lower = text.lower()
        
        for tech in technologies:
            if tech.lower() in text_lower:
                found_tech.append(tech)
        
        return found_tech[:10]  
